from pathlib import Path
import os
import json
import re

from openai import OpenAI
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from prompts import logic_prompt


md_path = Path(r"C:\agent\Data_Validation\output_markdown\가족보호자_임종기돌봄 설문지(최종)_심층인터뷰 추가_중.md")
exell_path = Path(r"C:\agent\Data_Validation\output_exell\가족보호자_임종기돌봄 설문지(최종)_심층인터뷰 추가_중.xlsx")

def to_excel(rows: list[dict]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "skip"

    col = 1  # A열부터

    for it in rows:
        start = (it.get("start") or "").strip()
        end = (it.get("end") or "").strip()
        value = it.get("value") or []

        # value를 리스트로 정규화
        values = value if isinstance(value, list) else [value]

        # value가 비어있으면 한 칸은 만들되 값은 공백
        if not values:
            values = [""]

        for v in values:
            # 1행: 로직명
            ws.cell(row=1, column=col, value="문항 스킵")
            # 2행: start,end
            ws.cell(row=2, column=col, value=f"{start},{end}".strip(","))
            # 3행: 선택번호 (숫자면 int로)
            if isinstance(v, int):
                ws.cell(row=3, column=col, value=v)
            else:
                s = str(v).strip()
                ws.cell(row=3, column=col, value=int(s) if s.isdigit() else s)

            col += 1

    for c in range(1, col):
        ws.column_dimensions[get_column_letter(c)].width = 18

    wb.save(exell_path)


def logic_llm():
    client = OpenAI(api_key="")
    md = md_path.read_text(encoding="utf-8")

    resp = client.responses.create(
        model="gpt-4.1-mini",
        input=[
            {"role": "system", "content": logic_prompt},
            {"role": "user", "content": md},
        ],
        temperature=0,
    )

    text = resp.output_text.strip()
    text = re.sub(r"^```json\s*|\s*```$", "", text, flags=re.IGNORECASE).strip()

    data = json.loads(text)

    to_excel(data)

    print("완료:", exell_path)


if __name__ == "__main__":
    logic_llm()